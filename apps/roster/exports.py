"""Download builders for roster exports."""
from collections import defaultdict
from datetime import timedelta
from html import escape
from io import BytesIO
from pathlib import Path

from django.utils import timezone


DETAIL_HEADERS = ("Employee", "Role", "Date", "Day", "Time", "Location")
DETAIL_FIELDS = ("employee", "role", "date", "day", "time", "location")


def export_rows(roster):
    assignments = roster.assignments.select_related("employee", "location").order_by("employee__name", "starts_at", "location__name")
    return [{"employee": item.employee.name, "role": item.employee.role, "date": timezone.localtime(item.starts_at).date(), "day": timezone.localtime(item.starts_at).strftime("%A"), "time": "%s-%s" % (timezone.localtime(item.starts_at).strftime("%H:%M"), timezone.localtime(item.ends_at).strftime("%H:%M")), "location": item.location.name, "hours": (item.ends_at - item.starts_at).total_seconds() / 3600} for item in assignments]


def rows_by_employee(rows):
    grouped = defaultdict(list)
    for row in rows:
        grouped[(row["employee"], row["role"])].append(row)
    return grouped


def _month_days(roster):
    day, days = roster.month, []
    while day.month == roster.month.month:
        days.append(day)
        day += timedelta(days=1)
    return days


def _hours_label(hours):
    return f"{hours:g} h"


def _workbook_sheet(workbook, title, selected_rows, roster):
    from openpyxl.styles import Font
    sheet = workbook.create_sheet(title)
    sheet.append([f"{roster.month:%B %Y} roster"])
    sheet.append(DETAIL_HEADERS)
    for cell in sheet[2]: cell.font = Font(bold=True)
    for row in selected_rows: sheet.append([row[field] if field != "date" else row[field].isoformat() for field in DETAIL_FIELDS])
    for column, width in zip("ABCDEF", (26, 18, 14, 14, 16, 24)): sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A3"


def excel_export(roster, scope):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from apps.people.models import Employee
    rows = export_rows(roster)
    workbook = Workbook(); workbook.remove(workbook.active)
    if scope != "total":
        used_titles = set()
        for title, selected_rows in [(employee, values) for (employee, _), values in rows_by_employee(rows).items()] or [("Roster", [])]:
            base, sheet_name, suffix = title[:31] or "Roster", title[:31] or "Roster", 2
            while sheet_name in used_titles: sheet_name, suffix = f"{base[:28]} ({suffix})", suffix + 1
            used_titles.add(sheet_name); _workbook_sheet(workbook, sheet_name, selected_rows, roster)
    else:
        sheet, days = workbook.create_sheet("Roster"), _month_days(roster)
        sheet.append([f"{roster.month:%B %Y} roster"])
        sheet.append(["Employee", "Role", *[day.day for day in days]])
        for cell in sheet[2]:
            cell.font = Font(bold=True); cell.alignment = Alignment(horizontal="center")
        scheduled = {(row["employee"], row["role"]) for row in rows}
        employees = [(employee.name, employee.role) for employee in Employee.objects.filter(active=True).order_by("name")] 
        employees.extend(sorted(scheduled - set(employees)))
        hours = defaultdict(float)
        for row in rows: hours[(row["employee"], row["role"], row["date"])] += row["hours"]
        for employee, role in employees:
            sheet.append([employee, role, *[_hours_label(hours[employee, role, day]) if hours[employee, role, day] else "" for day in days]])
        sheet.column_dimensions["A"].width, sheet.column_dimensions["B"].width = 26, 18
        for column in range(3, len(days) + 3): sheet.column_dimensions[get_column_letter(column)].width = 7
        for row in sheet.iter_rows(min_row=3, min_col=3):
            for cell in row: cell.alignment = Alignment(horizontal="center")
        weekend_fill = PatternFill("solid", fgColor="FFF2CC")
        for day_index, day in enumerate(days, start=3):
            if day.weekday() >= 5:
                for row_index in range(2, sheet.max_row + 1):
                    sheet.cell(row_index, day_index).fill = weekend_fill
            else:
                sheet.cell(2, day_index).fill = PatternFill("solid", fgColor="E9ECEF")
        sheet.freeze_panes = "C3"
    output = BytesIO(); workbook.save(output)
    return output.getvalue()


def _pdf_font():
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    for path in (Path("C:/Windows/Fonts/arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"), Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf")):
        if path.exists():
            pdfmetrics.registerFont(TTFont("RosterUnicode", str(path)))
            return "RosterUnicode"
    return "Helvetica"


def pdf_export(roster, scope):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    from .services import roster_calendar
    rows, output, font_name = export_rows(roster), BytesIO(), _pdf_font()
    employee_pages = scope == "employee_pages"
    document = SimpleDocTemplate(output, pagesize=A4 if employee_pages else landscape(A4), leftMargin=.65*cm if not employee_pages else 1.25*cm, rightMargin=.65*cm if not employee_pages else 1.25*cm, topMargin=.8*cm, bottomMargin=.8*cm, title=f"Roster {roster.month:%B %Y}")
    styles = getSampleStyleSheet(); styles["Heading1"].fontName = font_name
    cell_style = ParagraphStyle("calendar-cell", parent=styles["Normal"], fontName=font_name, fontSize=8, leading=9.5)
    story = []
    def table(title, selected_rows, widths):
        story.extend((Paragraph(escape(title), styles["Heading1"]), Spacer(1, .25*cm)))
        data = [DETAIL_HEADERS] + [[str(row[field]) if field == "date" else row[field] for field in DETAIL_FIELDS] for row in selected_rows]
        result = Table(data, colWidths=widths, repeatRows=1)
        result.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d6efd")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("FONTNAME", (0, 0), (-1, -1), font_name), ("GRID", (0, 0), (-1, -1), .25, colors.HexColor("#adb5bd")), ("FONTSIZE", (0, 0), (-1, -1), 8)]))
        story.append(result)
    if employee_pages:
        groups = list(rows_by_employee(rows).items())
        for index, ((employee, role), selected_rows) in enumerate(groups):
            table(f"{roster.month:%B %Y} - {employee if not role else f'{employee} - {role}'}", selected_rows, (3.2*cm, 2*cm, 2*cm, 1.7*cm, 2*cm, 3.2*cm))
            if index < len(groups)-1: story.append(PageBreak())
        if not groups: table(f"Roster - {roster.month:%B %Y}", [], (3.2*cm, 2*cm, 2*cm, 1.7*cm, 2*cm, 3.2*cm))
    else:
        headers = [Paragraph(day, cell_style) for day in ("Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun")]
        weeks = roster_calendar(roster)
        for week_index, week in enumerate(weeks):
            story.append(Paragraph(f"Roster - {roster.month:%B %Y}", styles["Heading1"]))
            data = [headers]
            calendar_row = []
            for cell in week:
                if not cell["in_month"]: calendar_row.append(""); continue
                parts = [f"<b>{cell['date'].day}</b>"]
                for shift in cell["shifts"]:
                    name = shift["template"].name if shift["template"] else "Custom shift"
                    assigned = "<br/>".join(escape(str(item.employee)) for item in shift["assignments"]) or "Unassigned"
                    parts.append(f"<b>{escape(name)}</b> {timezone.localtime(shift['starts_at']):%H:%M}-{timezone.localtime(shift['ends_at']):%H:%M}<br/>{escape(str(shift['location']))}<br/>{assigned}")
                calendar_row.append(Paragraph("<br/><br/>".join(parts), cell_style))
            data.append(calendar_row)
            calendar = Table(data, colWidths=[3.85*cm]*7, rowHeights=[.7*cm, 16.7*cm])
            calendar.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0d6efd")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), .35, colors.HexColor("#adb5bd")), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (5, 1), (6, -1), colors.HexColor("#fff8e1")), ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4), ("TOPPADDING", (0, 0), (-1, -1), 4)]))
            story.append(calendar)
            if week_index < len(weeks) - 1: story.append(PageBreak())
    document.build(story)
    return output.getvalue()
