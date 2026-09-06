from datetime import date
import pytest
from django.contrib.auth.models import User
from django.test import Client
from django.utils import timezone
from datetime import datetime
from apps.people.models import Employee, Location
from apps.roster.models import RosterVersion, ShiftAssignment
from io import BytesIO
from openpyxl import load_workbook

@pytest.mark.django_db
def test_export_screen_is_available_for_draft_and_published_rosters():
    roster = RosterVersion.objects.create(month=date(2026, 2, 1))
    client = Client(); client.force_login(User.objects.create_user("manager", password="test"))
    response = client.get(f"/roster/{roster.pk}/export/")
    assert response.status_code == 200
    assert b"employee_pages" in response.content and b"xlsx" in response.content and b"pdf" in response.content
    roster.status = RosterVersion.SUPERSEDED; roster.save(update_fields=["status"])
    assert client.get(f"/roster/{roster.pk}/export/").status_code == 404

@pytest.mark.django_db
def test_export_downloads_excel_and_pdf():
    roster = RosterVersion.objects.create(month=date(2026, 2, 1))
    employee, location = Employee.objects.create(name="Jane Manager"), Location.objects.create(name="Main office")
    starts = timezone.make_aware(datetime(2026, 2, 2, 9))
    ShiftAssignment.objects.create(roster_version=roster, employee=employee, location=location, starts_at=starts, ends_at=starts.replace(hour=17))
    client = Client(); client.force_login(User.objects.create_user("manager", password="test"))
    excel = client.get(f"/roster/{roster.pk}/export/?format=xlsx&scope=total")
    pdf = client.get(f"/roster/{roster.pk}/export/?format=pdf&scope=total")
    assert excel["Content-Type"].startswith("application/vnd.openxmlformats") and excel.content.startswith(b"PK")
    sheet = load_workbook(BytesIO(excel.content)).active
    assert [sheet.cell(2, column).value for column in range(1, 4)] == ["Employee", 1, 2]
    assert sheet.cell(3, 1).value == "Jane Manager"
    assert sheet.cell(3, 3).value == "8 h"
    assert sheet.cell(3, 30).value == "8 h"
    assert [sheet.cell(2, column).value for column in range(30, 33)] == ["HOURS", "PLAN", "COEF."]
    assert sheet.cell(2, 2).fill.fgColor.rgb == "00FFF2CC"
    assert sheet.cell(3, 3).fill.fgColor.rgb == "00E2F0D9"
    assert sheet.cell(4, 3).value == 1
    assert pdf["Content-Type"] == "application/pdf" and pdf.content.startswith(b"%PDF")
